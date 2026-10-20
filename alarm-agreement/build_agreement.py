"""Build the Total Alarm Pro installation + monitoring agreement (.docx).

Usage: python3 build_agreement.py OUTPUT.docx
"""
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

SELLER = "Sycamore Construction LLC dba Total Alarm Pro"
ADDRESS = "115 N College St, Mountain Home, AR 72653"
PHONE = "(870) 321-9063"
LICENSE = "CMPY.0003740"
BLANK = "________"

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.6)
    s.left_margin = s.right_margin = Inches(0.75)
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(3)

footer = doc.sections[0].footer.paragraphs[0]
footer.text = "DRAFT – FOR ATTORNEY REVIEW"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.runs[0].bold = True
footer.runs[0].font.size = Pt(9)


def para(text="", bold=False, size=None, align=None, space_after=None, italic=False):
    p = doc.add_paragraph()
    if text:
        add_runs(p, text, bold, size, italic)
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_runs(p, text, bold=False, size=None, italic=False):
    # **text** marks bold spans inside a normal paragraph.
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    return p


def heading(text):
    p = para(text, bold=True, size=11, space_after=2)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    return p


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def boxed(lines, bold=True, size=10):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t)
    cell = t.rows[0].cells[0]
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        add_runs(p, line, bold=bold, size=size)
    para(space_after=2)
    return t


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def grid(rows, widths=None, header=True, size=10):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    set_borders(t)
    if widths:
        t.autofit = False
        for c, w in enumerate(widths):
            t.columns[c].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.rows[r].cells[c]
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], val, bold=header and r == 0, size=size)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            if widths:
                cell.width = Inches(widths[c])
    para(space_after=2)
    return t


# ---------------------------------------------------------------- agreement
para("ALARM INSTALLATION AND MONITORING AGREEMENT", bold=True, size=14,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para("Form TAP-1", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

para(f"**{SELLER}** (“we” or “us”)", space_after=0)
para(f"{ADDRESS}  ·  {PHONE}", space_after=0)
para(f"Arkansas alarm company license (Class E) No. **{LICENSE}**", space_after=6)

boxed([
    "Regulated by: Division of Arkansas State Police",
    "1 State Police Plaza Drive",
    "Little Rock, Arkansas 72209",
    "501-618-8600",
], size=11)

heading("1. Customer and Site")
grid([
    ["Customer name", "", "Phone", ""],
    ["Service address", "", "Email", ""],
    ["Mailing address", "", "Site type", "☐ Residential   ☐ Commercial"],
], widths=[1.3, 2.6, 0.8, 2.3], header=False)

heading("2. Equipment and Services")
para("We will install the equipment listed in Exhibit A and monitor it. "
     "Exhibit A also describes how the system works. "
     "Monitoring may be performed by an independent third-party monitoring center.")

heading("3. Price and Payments")
grid([
    ["Installation price", f"${BLANK}", "Monthly monitoring fee", f"${BLANK} per month"],
    ["Total cash price", f"${BLANK}", "Down payment", f"${BLANK}"],
    ["Unpaid balance", f"${BLANK}", "", ""],
], widths=[1.4, 1.6, 1.8, 2.2], header=False)
para(f"Payment schedule: {BLANK} monthly payments of ${BLANK} due on the {BLANK} of each month.")

heading("4. Term and Renewal")
para(f"The initial term is {BLANK} months, starting {BLANK}.")
para("**After the initial term, this agreement renews month to month only. You have the right, "
     "without additional cost or penalty, to terminate this agreement at the end of the initial "
     "term or at the end of any renewal month by giving us written notice at least 30 days before "
     "the end of that term.**")

heading("5. Your Right to Cancel")
para("You may cancel within three business days as explained in the attached Notices of "
     "Cancellation. If you cancel, we will refund all of your payments within 10 days after we "
     "receive your notice.")

heading("6. Limited Warranty")
para(f"**What is covered.** Equipment we install and our installation work, against defects in "
     f"materials and workmanship. [CONFIRM AGAINST REV 8]")
para("**What we will do.** Repair or replace the defective equipment or work, at no charge. "
     "[CONFIRM AGAINST REV 8]")
para(f"**When it starts and how long it lasts.** Coverage starts on the date installation is "
     f"complete and lasts {BLANK} months.")
para(f"**How to make a claim.** Contact {SELLER}, {ADDRESS}, {PHONE}, and tell us what is wrong.")
para(f"**For residential customers, any implied warranties, including the implied warranties of "
     f"merchantability and fitness for a particular purpose, last only as long as this written "
     f"warranty ({BLANK} months).** Some States do not allow limitations on how long an implied "
     f"warranty lasts, so the above limitation may not apply to you.")
para("Section 7 limits incidental and consequential damages. Some States do not allow the "
     "exclusion or limitation of incidental or consequential damages, so the above limitation or "
     "exclusion may not apply to you.")
para("This warranty gives you specific legal rights, and you may also have other rights which "
     "vary from State to State.")

heading("7. Limit of Liability")
para("If we or the monitoring center are found liable to you for any loss, damage, or injury of "
     "any kind related to the equipment, its installation, or the monitoring service, including "
     "incidental and consequential damages, the most that we and the monitoring center together "
     "will owe you is $[____ SET BEFORE MONITORING-CENTER SUBMISSION]. This limit protects the "
     "monitoring center and its employees and agents the same way it protects us.")

heading("8. Mountain Home Alarm Permit")
para("If the service address is inside Mountain Home city limits, you must have an alarm user "
     "permit from the City of Mountain Home before we install the system.")

heading("9. Electronic Signature Consent (only if signed electronically)")
para("By initialing here ______, you agree to receive and sign this agreement, its exhibit, and "
     "the Notices of Cancellation electronically. This consent covers only those documents. You "
     "have the right to a paper copy of any of them; to get one, call or write us at the phone "
     "number or address above. You may withdraw this consent at any time by telling us in writing; "
     "after that we will send these documents on paper. To view and keep them you need a phone, "
     "tablet, or computer with internet access, a current web browser, a PDF reader, and a way to "
     "save or print.")

heading("10. Signatures")
para(f"Seller: {SELLER}, {ADDRESS}", space_after=4)
sig = doc.add_table(rows=2, cols=2)
set_borders(sig)
left = sig.rows[0].cells[0]
left.paragraphs[0].text = ""
add_runs(left.paragraphs[0],
         "You, the buyer, may cancel this transaction at any time prior to midnight of the third "
         "business day after the date of this transaction. See the attached notice of cancellation "
         "form for an explanation of this right.", bold=True, size=10)
right = sig.rows[0].cells[1]
right.paragraphs[0].text = ""
add_runs(right.paragraphs[0], "Customer signature:\n\n__________________________________\n"
         "Printed name: _____________________\nDate customer signs: _______________")
seller_cell = sig.rows[1].cells[1]
seller_cell.paragraphs[0].text = ""
add_runs(seller_cell.paragraphs[0], "For Total Alarm Pro:\n\n__________________________________\n"
         "Name and title: ___________________\nDate: _____________________________")
sig.rows[1].cells[0].paragraphs[0].text = ""
add_runs(sig.rows[1].cells[0].paragraphs[0],
         "Customer received a completed copy of this agreement and two copies of each Notice of "
         "Cancellation.  Customer initials: ______")
para(space_after=0)

# ---------------------------------------------------------------- exhibit A
page_break()
para("EXHIBIT A – EQUIPMENT AND SYSTEM DESCRIPTION", bold=True, size=12,
     align=WD_ALIGN_PARAGRAPH.CENTER)
grid([["Qty", "Equipment", "Location"]] + [["", "", ""] for _ in range(10)],
     widths=[0.6, 3.6, 2.8])
for label in ("Function (what the system detects and reports):",
              "Design (zones, devices, and how signals reach the monitoring center):",
              "Operation (how to arm and disarm, and what happens when an alarm goes off):",
              "Monitoring services provided:"):
    para(f"**{label}**", space_after=0)
    for _ in range(3):
        para("_" * 98, space_after=0)
    para(space_after=4)


# ---------------------------------------------------------------- notices
def cut_line():
    para("✂ - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - -",
         size=8, align=WD_ALIGN_PARAGRAPH.CENTER)


def federal_notice(copy_no):
    page_break()
    cut_line()
    b = dict(bold=True, size=10)
    para(f"Federal notice – copy {copy_no} of 2", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para("NOTICE OF CANCELLATION", align=WD_ALIGN_PARAGRAPH.CENTER, **b)
    para("______________________________\n(Date of transaction)", **b)
    for text in (
        "YOU MAY CANCEL THIS TRANSACTION, WITHOUT ANY PENALTY OR OBLIGATION, WITHIN THREE BUSINESS "
        "DAYS FROM THE ABOVE DATE.",
        "IF YOU CANCEL, ANY PROPERTY TRADED IN, ANY PAYMENTS MADE BY YOU UNDER THE CONTRACT OR SALE, "
        "AND ANY NEGOTIABLE INSTRUMENT EXECUTED BY YOU WILL BE RETURNED WITHIN TEN BUSINESS DAYS "
        "FOLLOWING RECEIPT BY THE SELLER OF YOUR CANCELLATION NOTICE, AND ANY SECURITY INTEREST "
        "ARISING OUT OF THE TRANSACTION WILL BE CANCELLED.",
        "IF YOU CANCEL, YOU MUST MAKE AVAILABLE TO THE SELLER AT YOUR RESIDENCE, IN SUBSTANTIALLY AS "
        "GOOD CONDITION AS WHEN RECEIVED, ANY GOODS DELIVERED TO YOU UNDER THIS CONTRACT OR SALE, OR "
        "YOU MAY, IF YOU WISH, COMPLY WITH THE INSTRUCTIONS OF THE SELLER REGARDING THE RETURN "
        "SHIPMENT OF THE GOODS AT THE SELLER'S EXPENSE AND RISK.",
        "IF YOU DO MAKE THE GOODS AVAILABLE TO THE SELLER AND THE SELLER DOES NOT PICK THEM UP WITHIN "
        "20 DAYS OF THE DATE OF YOUR NOTICE OF CANCELLATION, YOU MAY RETAIN OR DISPOSE OF THE GOODS "
        "WITHOUT ANY FURTHER OBLIGATION. IF YOU FAIL TO MAKE THE GOODS AVAILABLE TO THE SELLER, OR "
        "IF YOU AGREE TO RETURN THE GOODS TO THE SELLER AND FAIL TO DO SO, THEN YOU REMAIN LIABLE "
        "FOR PERFORMANCE OF ALL OBLIGATIONS UNDER THE CONTRACT.",
        f"TO CANCEL THIS TRANSACTION, MAIL OR DELIVER A SIGNED AND DATED COPY OF THIS CANCELLATION "
        f"NOTICE OR ANY OTHER WRITTEN NOTICE, OR SEND A TELEGRAM, TO {SELLER.upper()}, AT "
        f"{ADDRESS.upper()} NOT LATER THAN MIDNIGHT OF ______________________.",
        "I HEREBY CANCEL THIS TRANSACTION.",
    ):
        para(text, **b)
    para("______________________________\n(Date)", **b)
    para("______________________________\n(Buyer's signature)", **b)
    cut_line()


def arkansas_notice(copy_no):
    page_break()
    cut_line()
    para(f"Arkansas notice – copy {copy_no} of 2", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para("NOTICE OF CANCELLATION", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    boxed([
        "[INSERT HERE, WORD FOR WORD, THE FORM PRINTED IN ARK. CODE ANN. § 4-89-108(c).]",
        "",
        "Set it in 10-point bold. Fill every blank: seller's name "
        f"“{SELLER}”; seller's address “{ADDRESS}”; the date of the transaction; and the "
        "cancellation deadline (midnight of the third day after signing, not counting Sundays "
        "and holidays).",
        "",
        "UNVERIFIED: the official text could not be retrieved while drafting, so it has not been "
        "typed in from memory. Source: https://law.justia.com/codes/arkansas/title-4/subtitle-7/"
        "chapter-89/section-4-89-108/",
    ], bold=True, size=10)
    cut_line()


arkansas_notice(1)
arkansas_notice(2)
federal_notice(1)
federal_notice(2)

# ---------------------------------------------------------------- checklist
page_break()
para("COMPLIANCE CHECKLIST (not part of the agreement; remove before use)", bold=True, size=12,
     align=WD_ALIGN_PARAGRAPH.CENTER)
grid([
    ["#", "Requirement", "Where it appears"],
    ["1", "Company name and license no. (17 CAR § 350-509(a))", "Header"],
    ["2", "“Regulated by” block, bold, conspicuous (17 CAR § 350-509(c))", "Header box"],
    ["3", "Month-to-month renewal; bold termination right, 30 days' written notice "
          "(Ark. Code § 4-86-106)", "§ 4"],
    ["4", "Signed by both parties; seller name/address; date customer signs; copy to customer",
          "§ 10"],
    ["4", "Total cash price, down payment, unpaid balance, payment schedule", "§ 3"],
    ["4", "Description of equipment and services", "§ 2, Exhibit A"],
    ["4", "Two detachable Arkansas Notices of Cancellation (§ 4-89-108(c))",
          "Arkansas notices 1–2 (text still to be inserted)"],
    ["4", "Refund within 10 days (§ 4-89-109)", "§ 5"],
    ["5", "Bold cancellation statement next to customer signature (16 CFR 429.1(a))", "§ 10"],
    ["5", "Two federal Notices of Cancellation, 10-pt bold (16 CFR 429.1(b))",
          "Federal notices 1–2"],
    ["5", "No waiver of cancellation right; no confession of judgment",
          "Neither appears anywhere"],
    ["5", "Same language as sales presentation", "Agreement is in English; use only after "
          "English-language presentations"],
    ["5", "Separate or merged notices", "Kept separate (no merged form to review)"],
    ["6", "Implied warranties not disclaimed for residential; duration limited, prominent",
          "§ 6 (bold sentence)"],
    ["6", "Coverage, remedy, start and length, claim steps, three required sentences", "§ 6"],
    ["7", "ESIGN consent (15 U.S.C. 7001(c))", "§ 9"],
    ["8", "Mountain Home alarm user permit line (Ch. 4.40)", "§ 8"],
    ["9", "Written description of function, design, operation (§ 4.40.04(E))", "Exhibit A"],
    ["10", "Monitoring may be by independent third-party center", "§ 2"],
    ["11", "Liability limit that covers the monitoring center", "§ 7"],
    ["12", "One fixed version; blanks only for customer details, price, equipment",
          "Whole form (TAP-1). Set the § 7 dollar amount before submitting to UCC."],
], widths=[0.4, 4.0, 2.6], size=9)

page_break()
para("SUGGESTED, NOT ADDED", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Items that may be legally required but were not on the list, so they are not in the "
     "agreement. Each needs attorney confirmation.", italic=True)
for text in (
    "**Oral notice of the right to cancel at the time of sale.** A sales-practice duty, not "
    "contract text. 16 CFR 429.1(e), https://www.ecfr.gov/current/title-16/part-429; Ark. Code "
    "§ 4-89-108(d) (cooling-off period does not start until oral and written notice are given), "
    "https://law.justia.com/codes/arkansas/title-4/subtitle-7/chapter-89/section-4-89-108/.",
    "**“Limited Warranty” designation.** A written warranty on a consumer product over $10 must "
    "be titled “Full” or “Limited,” and a warranty that limits implied-warranty duration cannot "
    "be “Full.” Section 6 is headed “Limited Warranty” for this reason. 15 U.S.C. 2303, "
    "https://www.law.cornell.edu/uscode/text/15/2303.",
    "**Warranty available before sale.** The warranty text must be shown to the customer before "
    "the sale. 16 CFR 702, https://www.ecfr.gov/current/title-16/part-702.",
    "**ESIGN extras.** Before consent, the consumer must also be told how to update their "
    "electronic contact information and whether a paper copy costs anything, and must consent "
    "electronically in a way that shows they can open the documents. 15 U.S.C. "
    "7001(c)(1)(B)(iv), (c)(1)(C)(ii), https://www.law.cornell.edu/uscode/text/15/7001.",
    "**Paper notices when e-signing.** Arkansas requires the notices to be “easily detachable,” "
    "which suggests paper copies even for an e-signed contract. UNVERIFIED.",
    "**Refund timing conflict.** The federal notice promises refunds within 10 business days; "
    "§ 5 promises 10 days (the Arkansas rule, stricter). Confirm § 4-89-109 says 10 calendar days. "
    "UNVERIFIED.",
    "**Liability limit enforceability.** Whether Arkansas enforces a liability cap in a consumer "
    "contract for gross negligence or personal injury is a question for the attorney. UNVERIFIED.",
):
    p = para(text)
    p.paragraph_format.left_indent = Inches(0.2)

para("VERIFICATION NOTES", bold=True, size=11).paragraph_format.space_before = Pt(10)
for text in (
    "**Arkansas Notice of Cancellation text: UNVERIFIED, not inserted.** The official form could "
    "not be retrieved while drafting. Paste it from § 4-89-108(c) before use.",
    "**Federal Notice of Cancellation text:** typed from 16 CFR 429.1(b). Proofread word for word "
    "against https://www.ecfr.gov/current/title-16/section-429.1 before use.",
    "**17 CAR § 350-509:** the “Regulated by” statement and 501-618-8600 match Part 350 as "
    "published (current through Feb. 6, 2026), "
    "https://media.ark.org/dps/Code-of-Arkansas-Rules-%E2%80%93-Part-350.pdf. The "
    "license-number subsection (a) was not seen in full: UNVERIFIED.",
    "**Ark. Code § 4-86-106:** month-to-month renewal and 30-day written notice match the 2025 "
    "text, https://law.justia.com/codes/arkansas/title-4/subtitle-7/chapter-86/section-4-86-106/.",
    "**Ark. Code §§ 4-89-106, 4-89-109 and Mountain Home Ch. 4.40:** text not seen; drafted from "
    "the owner's summary. UNVERIFIED. City code: https://www.cityofmountainhome.com/municipal-code.",
    "**Rev 8 draft:** not available while drafting. Sections 6 “What is covered” and “What we "
    "will do” are marked [CONFIRM AGAINST REV 8].",
):
    p = para(text)
    p.paragraph_format.left_indent = Inches(0.2)

doc.save(sys.argv[1])
print("wrote", sys.argv[1])
